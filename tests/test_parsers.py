"""Tests for document parsers."""

import io

import pytest

from app.core.exceptions import DocumentParseError
from app.services.documents.parsers.docx import DocxParser
from app.services.documents.parsers.md import MarkdownParser
from app.services.documents.parsers.pdf import PdfParser
from app.services.documents.parsers.registry import build_default_registry


def test_markdown_parser():
    parser = MarkdownParser()
    content = b"# Hello Markdown\n\nThis is a test."
    result = parser.parse("test.md", content)

    assert result.media_type == "text/markdown"
    assert "Hello Markdown" in result.text
    assert result.metadata["parser"] == "markdown"
    assert result.metadata["filename"] == "test.md"


def test_markdown_parser_invalid_utf8():
    parser = MarkdownParser()
    content = b"\xff\xfe"
    with pytest.raises(DocumentParseError, match="decoded as UTF-8"):
        parser.parse("test.md", content)


def test_pdf_parser():
    # pypdf test document creation
    try:
        from pypdf import PdfWriter
    except ImportError:
        pytest.skip("pypdf not installed")

    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    # Adding text to a PDF programmatically without external fonts is complex in pypdf,
    # so we test metadata extraction and parser initialization only.
    writer.add_metadata({"/Title": "Test PDF", "/Author": "Test Author"})

    buf = io.BytesIO()
    writer.write(buf)
    content = buf.getvalue()

    parser = PdfParser()
    result = parser.parse("test.pdf", content)

    assert result.media_type == "application/pdf"
    assert result.metadata["parser"] == "pdf"
    assert result.metadata["filename"] == "test.pdf"
    assert result.metadata["pages"] == 1
    assert result.metadata["title"] == "Test PDF"
    assert result.metadata["author"] == "Test Author"
    # Text will be empty because we just added a blank page
    assert result.text == ""


def test_pdf_parser_invalid():
    parser = PdfParser()
    with pytest.raises(DocumentParseError, match="Could not read PDF"):
        parser.parse("test.pdf", b"not a pdf")


def test_docx_parser():
    try:
        import docx
    except ImportError:
        pytest.skip("python-docx not installed")

    doc = docx.Document()
    doc.core_properties.title = "Test DOCX"
    doc.core_properties.author = "Test Author"
    doc.add_paragraph("Hello DOCX")
    doc.add_paragraph("This is a test.")

    buf = io.BytesIO()
    doc.save(buf)
    content = buf.getvalue()

    parser = DocxParser()
    result = parser.parse("test.docx", content)

    assert (
        result.media_type
        == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    assert "Hello DOCX\n\nThis is a test." in result.text
    assert result.metadata["parser"] == "docx"
    assert result.metadata["filename"] == "test.docx"
    assert result.metadata["title"] == "Test DOCX"
    assert result.metadata["author"] == "Test Author"


def test_docx_parser_invalid():
    parser = DocxParser()
    with pytest.raises(DocumentParseError, match="Could not read DOCX"):
        parser.parse("test.docx", b"not a docx")


def test_registry_selects_correct_parsers():
    registry = build_default_registry()

    assert registry.select("test.md").name == "markdown"
    assert registry.select("test.markdown").name == "markdown"
    assert registry.select("test.pdf").name == "pdf"
    assert registry.select("test.docx").name == "docx"
    assert registry.select("test.txt").name == "txt"

    # Test by media type
    assert (
        registry.select("test.unknown", media_type="text/markdown").name == "markdown"
    )
    assert registry.select("test.unknown", media_type="application/pdf").name == "pdf"
